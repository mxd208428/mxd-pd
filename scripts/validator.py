#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""文件验证与安全保存模块"""

import os
from docx import Document


def validate_docx_file(file_path):
    """验证 Word 文档文件

    Args:
        file_path: 文件路径

    Returns:
        (is_valid, errors) — 是否有效，错误列表
    """
    errors = []

    # 检查文件是否存在
    if not os.path.exists(file_path):
        errors.append(f"文件不存在: {file_path}")
        return False, errors

    # 检查文件扩展名
    if not file_path.lower().endswith('.docx'):
        errors.append(f"文件格式不正确，需要 .docx 格式: {file_path}")
        return False, errors

    # 检查文件是否为空
    if os.path.getsize(file_path) == 0:
        errors.append(f"文件为空: {file_path}")
        return False, errors

    # 检查文件是否被占用
    try:
        with open(file_path, 'r+b'):
            pass
    except PermissionError:
        errors.append(f"文件被占用，请关闭 Word 后重试: {file_path}")
        return False, errors
    except OSError as e:
        errors.append(f"无法访问文件: {e}")
        return False, errors

    # 尝试打开文件
    try:
        doc = Document(file_path)
        if len(doc.paragraphs) == 0 and len(doc.tables) == 0:
            errors.append("文档内容为空")
            return False, errors
    except Exception as e:
        errors.append(f"文件损坏或无法打开: {e}")
        return False, errors

    return True, []


def safe_save_document(doc, output_path, overwrite=False):
    """安全保存文档

    Args:
        doc: Document 对象
        output_path: 输出路径
        overwrite: 是否覆盖已有文件

    Returns:
        (success, result) — 成功与否，成功时为实际保存路径，失败时为错误信息
    """
    # 检查输出目录是否存在
    output_dir = os.path.dirname(output_path)
    if output_dir and not os.path.exists(output_dir):
        try:
            os.makedirs(output_dir, exist_ok=True)
        except OSError as e:
            return False, f"无法创建输出目录: {e}"

    # 检查输出目录是否可写
    if output_dir:
        test_file = os.path.join(output_dir, '.write_test')
        try:
            with open(test_file, 'w') as f:
                f.write('test')
            os.remove(test_file)
        except (PermissionError, OSError):
            return False, f"输出目录不可写: {output_dir}"

    # 如果不覆盖，自动生成新文件名
    if not overwrite and os.path.exists(output_path):
        base, ext = os.path.splitext(output_path)
        counter = 1
        while os.path.exists(f"{base}_{counter}{ext}"):
            counter += 1
        output_path = f"{base}_{counter}{ext}"

    # 保存文件
    try:
        doc.save(output_path)
        return True, output_path
    except PermissionError:
        return False, "文件被占用，请关闭 Word 后重试"
    except Exception as e:
        return False, f"保存失败: {e}"


def generate_output_path(input_path, suffix="_已排版"):
    """根据输入路径生成默认输出路径

    Args:
        input_path: 输入文件路径
        suffix: 后缀名

    Returns:
        输出文件路径
    """
    name, ext = os.path.splitext(input_path)
    return f"{name}{suffix}{ext}"
