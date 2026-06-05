#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""格式检查模块（只报告问题，不修改文档）"""

from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from .formatter import (
    detect_heading_by_style,
    is_figure_caption,
    is_table_caption,
    is_reference,
    has_image,
)
from .captions import check_all_captions


def _get_run_font_info(run):
    """获取 run 的字体信息

    Returns:
        dict: {'font_cn': str, 'font_en': str, 'size_pt': float, 'bold': bool, 'color': str}
    """
    info = {
        'font_cn': '',
        'font_en': run.font.name or '',
        'size_pt': run.font.size.pt if run.font.size else 0,
        'bold': run.bold or False,
        'color': '',
    }

    # 获取中文字体
    rPr = run._element.find(qn('w:rPr'))
    if rPr is not None:
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            info['font_cn'] = rFonts.get(qn('w:eastAsia'), '')

    # 获取颜色
    if run.font.color and run.font.color.rgb:
        info['color'] = str(run.font.color.rgb)

    return info


def _get_paragraph_info(paragraph):
    """获取段落的格式信息

    Returns:
        dict
    """
    pf = paragraph.paragraph_format
    info = {
        'alignment': str(pf.alignment) if pf.alignment else 'None',
        'line_spacing': pf.line_spacing.pt if pf.line_spacing else 0,
        'first_indent': pf.first_line_indent.pt if pf.first_line_indent else 0,
        'space_before': pf.space_before.pt if pf.space_before else 0,
        'space_after': pf.space_after.pt if pf.space_after else 0,
        'runs': [],
    }

    for run in paragraph.runs:
        info['runs'].append(_get_run_font_info(run))

    return info


def _check_paragraph_format(para_info, expected_config, para_label):
    """检查段落格式是否符合规范

    Args:
        para_info: _get_paragraph_info 返回的字典
        expected_config: 期望的配置字典
        para_label: 段落描述（用于报告）

    Returns:
        问题列表
    """
    issues = []

    # 检查行距
    if expected_config.get('line_spacing'):
        expected_ls = expected_config['line_spacing']
        if hasattr(expected_ls, 'pt'):
            expected_ls = expected_ls.pt
        actual_ls = para_info['line_spacing']
        if actual_ls > 0 and abs(actual_ls - expected_ls) > 0.5:
            issues.append(f"{para_label}: 行距应为{expected_ls}pt，实际为{actual_ls}pt")

    # 检查首行缩进
    if 'first_line_indent' in expected_config:
        expected_fi = expected_config['first_line_indent']
        if hasattr(expected_fi, 'pt'):
            expected_fi = expected_fi.pt
        actual_fi = para_info['first_indent']
        if expected_fi > 0 and abs(actual_fi - expected_fi) > 1:
            issues.append(f"{para_label}: 首行缩进应为{expected_fi}pt，实际为{actual_fi}pt")

    # 检查 run 字体
    for j, run_info in enumerate(para_info['runs']):
        run_label = f"{para_label} 第{j+1}个文字块"

        # 检查字号
        if expected_config.get('size'):
            expected_size = expected_config['size']
            if hasattr(expected_size, 'pt'):
                expected_size = expected_size.pt
            actual_size = run_info['size_pt']
            if actual_size > 0 and abs(actual_size - expected_size) > 0.5:
                issues.append(f"{run_label}: 字号应为{expected_size}pt，实际为{actual_size}pt")

        # 检查中文字体
        if expected_config.get('font_cn') and run_info['font_cn']:
            if run_info['font_cn'] != expected_config['font_cn']:
                issues.append(f"{run_label}: 中文字体应为{expected_config['font_cn']}，实际为{run_info['font_cn']}")

        # 检查颜色
        if run_info['color'] and run_info['color'] != '000000':
            issues.append(f"{run_label}: 文字颜色不是黑色（{run_info['color']}）")

    return issues


def check_format(doc, config):
    """对照规范逐项检查文档格式

    Args:
        doc: Document 对象
        config: 已解析的配置字典

    Returns:
        dict: {
            'issues': 问题列表,
            'summary': 统计摘要,
            'caption_check': 图表序号检查结果,
        }
    """
    issues = []
    stats = {
        'headings': 0,
        'body': 0,
        'reference': 0,
        'figure_caption': 0,
        'table_caption': 0,
        'tables': 0,
        'images': 0,
        'non_black_text': 0,
    }

    in_references = False

    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue

        para_info = _get_paragraph_info(paragraph)
        para_label = f"第{i+1}段 \"{text[:30]}{'...' if len(text) > 30 else ''}\""

        # 检测标题
        style_level = detect_heading_by_style(paragraph)
        if style_level > 0:
            cfg_key = f"heading{style_level}"
            if cfg_key in config:
                para_issues = _check_paragraph_format(para_info, config[cfg_key], para_label)
                issues.extend(para_issues)
            stats['headings'] += 1

            if style_level == 1 and "参考文献" in text:
                in_references = True
            continue

        # 参考文献
        if in_references and is_reference(text):
            para_issues = _check_paragraph_format(para_info, config["reference"], para_label)
            issues.extend(para_issues)
            stats['reference'] += 1
            continue

        # 图题
        if is_figure_caption(text):
            para_issues = _check_paragraph_format(para_info, config["figure_caption"], para_label)
            issues.extend(para_issues)
            stats['figure_caption'] += 1
            continue

        # 表题
        if is_table_caption(text):
            para_issues = _check_paragraph_format(para_info, config["table_caption"], para_label)
            issues.extend(para_issues)
            stats['table_caption'] += 1
            continue

        # 摘要标题
        if "摘要" in text and len(text) < 10:
            para_issues = _check_paragraph_format(para_info, config["abstract_title"], para_label)
            issues.extend(para_issues)
            continue

        # 正文
        para_issues = _check_paragraph_format(para_info, config["body"], para_label)
        issues.extend(para_issues)
        stats['body'] += 1

        # 检查图片
        if has_image(paragraph):
            stats['images'] += 1

    # 检查表格
    for table in doc.tables:
        stats['tables'] += 1
        table_cfg = config.get("table", config["body"])
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                for para in cell.paragraphs:
                    for run in para.runs:
                        run_info = _get_run_font_info(run)
                        if run_info['color'] and run_info['color'] != '000000':
                            issues.append(f"表格第{row_idx+1}行第{col_idx+1}列: 文字颜色不是黑色")

    # 检查全文非黑色文字
    for para in doc.paragraphs:
        for run in para.runs:
            run_info = _get_run_font_info(run)
            if run_info['color'] and run_info['color'] != '000000':
                stats['non_black_text'] += 1

    # 图表序号检查
    caption_check = check_all_captions(doc)
    if caption_check['missing_figures']:
        issues.append(f"有 {len(caption_check['missing_figures'])} 个图片缺少图序图题")
    if caption_check['missing_tables']:
        issues.append(f"有 {len(caption_check['missing_tables'])} 个表格缺少表序表题")
    if caption_check['figure_seq_errors']:
        issues.append(f"图序号不连续: {', '.join(caption_check['figure_seq_errors'])}")
    if caption_check['table_seq_errors']:
        issues.append(f"表序号不连续: {', '.join(caption_check['table_seq_errors'])}")

    return {
        'issues': issues,
        'summary': stats,
        'caption_check': caption_check,
    }


def format_check_report(check_result):
    """将检查结果格式化为可读文本

    Args:
        check_result: check_format 返回的字典

    Returns:
        格式化的报告文本
    """
    lines = []
    lines.append("=" * 50)
    lines.append("格式检查报告")
    lines.append("=" * 50)

    # 统计摘要
    stats = check_result['summary']
    lines.append("")
    lines.append("【文档统计】")
    lines.append(f"  标题: {stats['headings']} 个")
    lines.append(f"  正文段落: {stats['body']} 个")
    lines.append(f"  参考文献: {stats['reference']} 条")
    lines.append(f"  图题: {stats['figure_caption']} 个")
    lines.append(f"  表题: {stats['table_caption']} 个")
    lines.append(f"  表格: {stats['tables']} 个")
    lines.append(f"  图片: {stats['images']} 个")
    if stats['non_black_text'] > 0:
        lines.append(f"  ⚠ 非黑色文字: {stats['non_black_text']} 处")

    # 图表序号
    caption = check_result['caption_check']
    lines.append("")
    lines.append("【图表序号】")
    lines.append(f"  图序号: {caption['figure_nums'] if caption['figure_nums'] else '无'}")
    lines.append(f"  表序号: {caption['table_nums'] if caption['table_nums'] else '无'}")

    # 问题清单
    issues = check_result['issues']
    lines.append("")
    lines.append(f"【问题清单】共 {len(issues)} 个问题")
    lines.append("-" * 50)

    if not issues:
        lines.append("  ✅ 未发现格式问题")
    else:
        for i, issue in enumerate(issues, 1):
            lines.append(f"  {i}. {issue}")

    lines.append("=" * 50)
    return "\n".join(lines)
